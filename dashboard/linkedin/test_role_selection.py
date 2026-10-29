import ast
import csv
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from role_selection import select_roles, should_split_states

ROOT = Path(__file__).resolve().parents[2]


def catalog():
    tree = ast.parse((ROOT / 'dashboard/linkedin/linkedin_68_roles.py').read_text())
    return next(ast.literal_eval(node.value) for node in tree.body
                if isinstance(node, ast.AnnAssign) and isinstance(node.target, ast.Name)
                and node.target.id == 'ROLES')


class RoleSelectionTests(unittest.TestCase):
    def test_three_workflows_cover_all_80_roles_without_overlap(self):
        roles = catalog()
        self.assertEqual(len(roles), 80)
        combined = []
        for start, end in ((1, 23), (24, 45), (46, 80)):
            selected = roles[start-1:end]
            shards = [select_roles(roles, filters=selected, start_idx=start, end_idx=end,
                                   shard=f'{part}/8') for part in range(1, 9)]
            flattened = [role for shard in shards for role in shard]
            self.assertEqual(flattened, selected)
            combined.extend(flattened)
        self.assertEqual(combined, roles)
        self.assertEqual(len({role.casefold() for role in combined}), 80)

    def test_default_selection_is_balanced_and_scraped_once(self):
        roles = catalog()
        for total in (4, 8, 12, 16):
            shards = [select_roles(roles, filters=roles[:23], shard=f'{part}/{total}')
                      for part in range(1, total + 1)]
            flattened = [role for shard in shards for role in shard]
            self.assertEqual(flattened, roles[:23])
            self.assertEqual(len(flattened), len(set(flattened)))
            self.assertLessEqual(max(map(len, shards)) - min(map(len, shards)), 1)
        self.assertEqual([len(select_roles(roles, filters=roles[:23], shard=f'{p}/8'))
                          for p in range(1, 9)], [3, 3, 3, 3, 3, 3, 3, 2])

    def test_overlapping_substrings_do_not_repeat_roles(self):
        roles = ['Senior Buyer', 'Buyer', 'Engineer']
        self.assertEqual(select_roles(roles, filters=['buyer', ' Senior BUYER ', 'buyer']), roles[:2])

    def test_empty_filter_uses_configured_index_range(self):
        roles = catalog()
        shards = [select_roles(roles, start_idx=24, end_idx=45, shard=f'{p}/8') for p in range(1, 9)]
        self.assertEqual([r for s in shards for r in s], roles[23:45])

    def test_more_runners_than_roles_and_invalid_input(self):
        self.assertEqual(select_roles(['Buyer'], shard='8/8'), [])
        for shard in ('0/8', '9/8', '1/0', 'bad', '1/2/3'):
            with self.assertRaises(ValueError):
                select_roles(['Buyer'], shard=shard)
        with self.assertRaises(ValueError):
            select_roles(['Buyer'], filters=['unknown role'])

    def test_keyword_queries_are_partitioned_exactly_once(self):
        queries = ['procurement', 'commodity', 'category', 'buyer', 'purchasing', 'Sourcing']
        shards = [select_roles(queries, shard=f'{part}/6') for part in range(1, 7)]
        self.assertEqual(shards, [[query] for query in queries])

    def test_every_keyword_can_split_and_all_mode_overrides_old_restrictions(self):
        queries = ['procurement', 'commodity', 'category', 'buyer', 'purchasing', 'Sourcing']
        for query in queries:
            with self.subTest(query=query):
                self.assertTrue(should_split_states(query, 'United States', True))
                self.assertFalse(should_split_states(query, 'United States', False))
                self.assertTrue(should_split_states(query, 'United States', False, mode='all'))
                self.assertTrue(should_split_states(query, 'United States', True, mode='all',
                                                    terms={'procurement', 'category', 'commodity'}))
                self.assertFalse(should_split_states(query, 'Canada', True, mode='all'))

    def test_state_mode_applies_to_all_catalog_roles_and_keyword_subsets(self):
        roles = catalog()
        self.assertTrue(all(should_split_states(role, 'USA', True) for role in roles))
        self.assertTrue(all(should_split_states(role, 'US', False, mode='all') for role in roles))
        queries = ['buyer', 'purchasing', 'Sourcing']
        shards = [select_roles(queries, shard=f'{part}/6') for part in range(1, 7)]
        self.assertEqual([role for shard in shards for role in shard], queries)
        self.assertEqual(sum(not shard for shard in shards), 3)
        with self.assertRaises(ValueError):
            should_split_states('buyer', 'USA', True, mode='disabled')

    def test_workflows_pass_role_selection_and_state_split_mode(self):
        import yaml
        def workflow(name):
            return yaml.load((ROOT / '.github/workflows' / name).read_text(), Loader=yaml.BaseLoader)
        procurement = workflow('linkedin_procurement_keywords.yml')
        inputs = procurement['on']['workflow_dispatch']['inputs']
        self.assertEqual(inputs['roles']['default'].split(','),
                         ['procurement', 'commodity', 'category', 'buyer', 'purchasing', 'Sourcing'])
        self.assertEqual(inputs['state_split_mode']['default'], 'all')
        self.assertIn('inputs.roles', procurement['jobs']['scrape']['with']['search_terms'])
        self.assertNotIn('state_split_terms', procurement['jobs']['scrape']['with'])
        for name in ['linkedin_roles_slice.yml', 'linkedin_roles_24_45.yml', 'linkedin_roles_46_68.yml']:
            item = workflow(name)
            mode = item['on']['workflow_dispatch']['inputs']['state_split_mode']
            self.assertEqual(mode['options'], ['auto', 'all'])
            if name != 'linkedin_roles_slice.yml':
                self.assertIn('inputs.state_split_mode', item['jobs']['scrape']['with']['state_split_mode'])
        shared = workflow('linkedin_roles_slice.yml')
        self.assertIn('state_split_mode', shared['on']['workflow_call']['inputs'])
        self.assertIn('--state-split-mode "$STATE_SPLIT_MODE"',
                      (ROOT / '.github/workflows/linkedin_roles_slice.yml').read_text())

    def test_screenshot_additions_are_unique_and_in_the_last_workflow(self):
        import yaml
        additions = ['Director of Procurement', 'Procurement Analyst', 'Commodity Trader',
                     'Director of Purchasing', 'Category Buyer', 'Strategic Sourcing Analyst',
                     'Media Buyer', 'Commodity Broker', 'Sourcing Associate', 'Strategic Sourcing Lead',
                     'Sourcing Director', 'Global Sourcing Manager']
        roles = catalog()
        self.assertEqual(roles[68:], additions)
        self.assertEqual([roles[i-1] for i in [1, 23, 45, 68]],
                         ['Senior Commodity Manager', 'Senior Procurement Consultant', 'Procurement Buyer', 'Purchasing Agent'])
        workflow = yaml.load((ROOT / '.github/workflows/linkedin_roles_46_68.yml').read_text(), Loader=yaml.BaseLoader)
        self.assertEqual(workflow['jobs']['scrape']['with']['role_end'], '80')
        filters = workflow['on']['workflow_dispatch']['inputs']['roles']['default'].split(',')
        self.assertEqual(select_roles(roles, filters=filters, start_idx=46, end_idx=80), roles[45:])

    def test_workflow_exports_one_all_sheet(self):
        from openpyxl import load_workbook
        workflow = (ROOT / '.github/workflows/linkedin_roles_slice.yml').read_text()
        # Execute the real merge/export code using saved-row fixtures, never a scraper.
        code = workflow.split("python3 - <<'PYEOF'\n", 1)[1].split('          PYEOF', 1)[0]
        import textwrap
        code = textwrap.dedent(code)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'dashboard').mkdir()
            for shard, rows in enumerate(([
                {'id': 'one', 'job_title': 'Buyer', 'description': 'Salary: $90-110K per year.',
                 'job_url': 'https://www.linkedin.com/jobs/view/4999999999', 'reposted': True,
                 'reposted_source': 'linkedin_label', 'reposted_confidence': 'explicit'},
            ], [
                {'id': 'one', 'job_title': 'Buyer', 'description': 'Duplicate row', 'job_url': 'https://example.com/1'},
                {'id': 'two', 'job_title': 'Manager', 'description': 'Manage sourcing.\nRequired: SQL.', 'job_url': 'https://example.com/2'},
            ]), 1):
                folder = root / f'scraped_data/shard-{shard}'
                folder.mkdir(parents=True)
                with (folder / f'linkedin_shard_{shard}-8_jobs.csv').open('w', newline='') as f:
                    writer = csv.DictWriter(f, fieldnames=list(rows[0]))
                    writer.writeheader()
                    writer.writerows(rows)
            cwd = Path.cwd()
            try:
                with patch.dict(os.environ, {'GITHUB_WORKSPACE': str(root), 'SCRAPER_DIR': 'dashboard',
                                             'ROLE_START': '1', 'ROLE_END': '23'}):
                    with patch.object(sys, 'path', [str(ROOT / 'dashboard'), str(ROOT / 'dashboard/db'), *sys.path]):
                        exec(compile(code, 'workflow-merge', 'exec'), {})
            finally:
                os.chdir(cwd)
            files = list((root / 'dashboard/deliverables').iterdir())
            self.assertEqual([f.name for f in files], ['jobs_roles_1_to_23.xlsx'])
            book = load_workbook(files[0], read_only=True)
            self.assertEqual(book.sheetnames, ['All'])
            rows = list(book['All'].values)
            self.assertEqual(len(rows), 3)
            id_column = rows[0].index('id')
            self.assertEqual({r[id_column] for r in rows[1:]}, {'one', 'two'})
            by_id = {r[id_column]: r for r in rows[1:]}
            self.assertEqual(by_id['one'][rows[0].index('reposted_source')], 'linkedin_label')
            self.assertEqual(by_id['one'][rows[0].index('repost_status')], 'Reposted')
            self.assertEqual(by_id['two'][rows[0].index('repost_status')], 'Unknown')
            book.close()


if __name__ == '__main__':
    unittest.main()
