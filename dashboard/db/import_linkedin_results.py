"""Validate completed LinkedIn All workbooks and idempotently upsert their jobs."""
import argparse
import hashlib
import json
import os
from pathlib import Path
import sys
from urllib.parse import urlsplit

from openpyxl import load_workbook
import psycopg

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from db.neon_upsert import map_job, save_jobs
from db.linkedin_reposts import linkedin_id


def read_workbook(folder):
    files = list(folder.rglob('*.xlsx'))
    if len(files) != 1:
        raise ValueError('Expected exactly one completed workbook per source run')
    path = files[0]
    book = load_workbook(path, read_only=True, data_only=True)
    if 'All' not in book.sheetnames:
        raise ValueError('Source workbook must contain the All sheet')
    values = book['All'].values
    header = list(next(values))
    if not {'id', 'job_title', 'job_url', 'description'}.issubset(header):
        raise ValueError('Source workbook is missing required job fields')
    rows = []
    for values_row in values:
        row = dict(zip(header, values_row))
        if not any(v is not None for v in values_row):
            continue
        host = (urlsplit(str(row.get('job_url') or '')).hostname or '').lower()
        if host not in ('linkedin.com', 'www.linkedin.com') or not linkedin_id(row.get('job_url')) or not row.get('id') or not row.get('job_title'):
            raise ValueError('Invalid LinkedIn job identity; refusing partial or guessed import')
        row['source_board'] = 'LinkedIn'
        extraction = row.get('extraction')
        if isinstance(extraction, str) and extraction:
            row['extraction'] = json.loads(extraction)
        rows.append(row)
    book.close()
    if not rows:
        raise ValueError('Source workbook contains no jobs')
    with path.open('rb') as f:
        sha = hashlib.file_digest(f, 'sha256').hexdigest()
    return rows, {'file': path.name, 'sha256': sha, 'rows': len(rows)}


def combine(sources):
    unique = {}
    for row in sources:
        identity = str(linkedin_id(row['job_url']))
        previous = unique.get(identity)
        if previous is None:
            unique[identity] = row
            continue
        # Keep the richest retrieved description, and preserve explicit repost
        # evidence even if it came from the other role range's copy.
        best, other = (row, previous) if len(row.get('description') or '') > len(previous.get('description') or '') else (previous, row)
        if other.get('reposted_confidence') == 'explicit' and best.get('reposted_confidence') != 'explicit':
            best = dict(best)
            for key in ('reposted', 'reposted_source', 'reposted_confidence', 'reposted_evidence'):
                best[key] = other.get(key)
        unique[identity] = best
    return unique


def assign_ids(unique, dbs):
    """Keep legacy IDs only when they belong to this exact posting, not a title lookalike."""
    candidates, owners = {}, {}
    for alias, url in dbs:
        with psycopg.connect(url, connect_timeout=30) as conn:
            conn.read_only = True
            conn.execute("SET statement_timeout = '2min'")
            for stored_id, stored_url in conn.execute("SELECT id, job_url FROM public.jobs WHERE lower(source_board)='linkedin'"):
                job_id = linkedin_id(stored_url)
                if not job_id:
                    continue
                key = str(job_id)
                owners.setdefault(str(stored_id), set()).add(key)
                if key in unique:
                    candidates.setdefault(key, []).append(str(stored_id))
    reassigned = 0
    for key, row in unique.items():
        existing = next((identity for identity in candidates.get(key, []) if owners[identity] == {key}), None)
        resolved = existing or ('linkedin:' + key)
        if resolved in owners and owners[resolved] != {key}:
            raise ValueError('A proposed posting identity already belongs to a different job')
        reassigned += str(row['id']) != resolved
        row['id'] = resolved
    if len({row['id'] for row in unique.values()}) != len(unique):
        raise ValueError('Distinct job postings must have distinct database IDs')
    return reassigned


def probe(dbs, ids):
    found, result = set(), []
    for alias, url in dbs:
        with psycopg.connect(url, connect_timeout=30) as conn:
            conn.read_only = True
            conn.execute("SET statement_timeout = '2min'")
            total, linkedin = conn.execute("SELECT count(*), count(*) FILTER (WHERE lower(source_board)='linkedin') FROM public.jobs").fetchone()
            matches = set()
            for stored_id, source in conn.execute('SELECT id, source_board FROM public.jobs WHERE id::text=ANY(%s::text[])', (ids,)):
                if (source or '').lower() != 'linkedin':
                    raise ValueError('An import identity collides with a different job board')
                matches.add(str(stored_id))
            found.update(matches)
            posts = None
            if conn.execute("SELECT to_regclass('public.linkedin_posts')").fetchone()[0]:
                posts = conn.execute('SELECT count(*) FROM public.linkedin_posts').fetchone()[0]
            result.append({'database': alias, 'total_jobs': total, 'linkedin_jobs': linkedin,
                           'source_ids_present': len(matches), 'hiring_posts_if_present': posts})
    return found, result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=Path('linkedin-import-input'))
    parser.add_argument('--output', type=Path, default=Path('linkedin-import-report.json'))
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    a, info_a = read_workbook(args.input / 'roles-1-23')
    b, info_b = read_workbook(args.input / 'roles-24-45')
    unique = combine(a + b)
    dbs = [(key, os.environ[key]) for key in ('NEON_DATABASE_URL', 'JOBS_DATABASE_URL_2',
                                             'JOBS_DATABASE_URL_3', 'JOBS_DATABASE_URL_4') if os.getenv(key)]
    if not dbs or dbs[0][0] != 'NEON_DATABASE_URL':
        raise RuntimeError('Main jobs database must be configured')
    reassigned = assign_ids(unique, dbs)
    ids = [row['id'] for row in unique.values()]
    existing, before = probe(dbs, ids)
    report = {'sources': [info_a, info_b], 'unique_jobs': len(unique),
              'duplicates_merged': len(a) + len(b) - len(unique),
              'already_in_database': len(existing), 'new_jobs': len(unique) - len(existing),
              'ids_resolved_by_posting_url': reassigned,
              'before': before, 'apply': args.apply, 'expire_missing': False,
              'storage_deletion_enabled': False, 'rows_written': 0}
    args.output.write_text(json.dumps(report, indent=2))
    if args.apply:
        # The user authorized uploading these results, not deleting more jobs.
        os.environ['JOB_RETENTION_DRY_RUN'] = 'true'
        for key in list(os.environ):
            if key == 'JOBS_DB_MAX_MB' or key.startswith('JOBS_DB_MAX_MB_'):
                os.environ.pop(key)
        report['rows_written'] = save_jobs(list(unique.values()), expire_missing=False)
        after_ids, report['after'] = probe(dbs, ids)
        report['all_source_ids_verified'] = set(ids).issubset(after_ids)
        args.output.write_text(json.dumps(report, indent=2))
        if report['rows_written'] != len(unique) or not report['all_source_ids_verified']:
            raise RuntimeError('Not every source job was verified after upload; successful upserts remain saved')
    print(json.dumps(report, indent=2), flush=True)


if __name__ == '__main__':
    main()
