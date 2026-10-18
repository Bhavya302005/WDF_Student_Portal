from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from io import BytesIO
import unittest

from docx import Document

from matching.config import MatchingSettings
from matching.factory import local_intake_stack
from matching.intake import IntakeRecommendationService
from matching.resume import CandidateIntakeService, ResumeParser, resume_text_from_bytes
from matching.schemas import ResumeRecommendationRequest
from matching.taxonomy import Taxonomy


RESUME = """Amin Example
PROFESSIONAL EXPERIENCE
Senior Backend Software Engineer, Acme
Jan 2020 - Present
Built Python and FastAPI services backed by PostgreSQL on AWS.

Software Engineering Intern, Example Labs
Jun 2019 - Dec 2019
Developed Python automation.

PROJECTS
Portfolio application using React and Node.js.

SKILLS
Python, FastAPI, PostgreSQL, AWS, React, Node.js

EDUCATION
Bachelor's in Computer Science

CERTIFICATIONS
AWS Certified Solutions Architect
"""


def required_answers(**overrides):
    answers = {
        "target_roles": ["Backend Engineer"],
        "confirmed_skills": ["Python", "FastAPI", "PostgreSQL", "AWS"],
        "professional_years": 4,
        "preferred_locations": ["Bengaluru"],
        "work_modes": ["hybrid", "remote"],
        "employment_types": ["full time"],
        "hard_location_constraint": False,
        "career_level": "mid",
        "work_authorization": "authorized without sponsorship",
        "requires_sponsorship": False,
    }
    answers.update(overrides)
    return answers


class CountingResumeParser(ResumeParser):
    def __init__(self):
        super().__init__(Taxonomy(), now=datetime(2024, 1, 1))
        self.calls = 0

    def parse(self, user_id, text):
        self.calls += 1
        return super().parse(user_id, text)


class ResumeIntakeTests(unittest.TestCase):
    def setUp(self):
        self.parser = ResumeParser(Taxonomy(), now=datetime(2024, 1, 1))
        self.intake = CandidateIntakeService(self.parser)

    def test_parser_keeps_professional_internship_and_project_evidence_separate(self):
        result = self.intake.parse_text("candidate-a", RESUME)

        self.assertIn("backend engineering", {role.casefold() for role in result.profile.roles})
        self.assertEqual(4.1, result.profile.professional_years)
        self.assertEqual(0.6, result.profile.internship_years)
        self.assertIn("fastapi", result.profile.professional_skills)
        self.assertIn("react", result.profile.project_skills)
        self.assertNotIn("react", result.profile.professional_skills)
        self.assertIn("Bachelor's in Computer Science", result.profile.education)
        self.assertTrue(any("AWS Certified" in value for value in result.profile.certifications))
        self.assertTrue(result.evidence["experience_dates"])
        self.assertTrue(all(question.id for question in result.questions))

    def test_editable_answers_override_parsed_values_and_validate_required_fields(self):
        parsed = self.intake.parse_text("candidate-a", RESUME)
        with self.assertRaisesRegex(ValueError, "must be confirmed"):
            self.intake.apply_answers(parsed, {})

        confirmed = self.intake.apply_answers(parsed, required_answers(
            target_roles=["Data Engineer"],
            confirmed_skills=["Python", "PostgreSQL"],
            professional_years=3.5,
            work_modes=["Remote"],
            hard_location_constraint="false",
        ))

        self.assertEqual(["Data Engineer"], confirmed.profile.roles)
        self.assertEqual(["python", "postgresql"], confirmed.profile.skills)
        self.assertEqual(3.5, confirmed.profile.professional_years)
        self.assertEqual(["remote"], confirmed.profile.work_mode_preferences)
        self.assertTrue(confirmed.profile.remote_preference)
        self.assertFalse(confirmed.profile.hard_location_constraint)
        self.assertEqual("authorized without sponsorship", confirmed.profile.work_authorization)
        self.assertFalse(confirmed.profile.requires_sponsorship)
        self.assertTrue(confirmed.profile.strict_role_constraint)
        self.assertTrue(confirmed.profile.strict_work_mode_constraint)
        self.assertEqual(30, confirmed.profile.max_job_age_days)
        self.assertTrue(all(question.confirmed for question in confirmed.questions if question.required))

        blank = {**required_answers(), "target_roles": []}
        with self.assertRaisesRegex(ValueError, "Which roles"):
            self.intake.apply_answers(parsed, blank)

    def test_same_resume_is_parsed_once_and_reused_safely_across_users(self):
        parser = CountingResumeParser()
        intake = CandidateIntakeService(parser)

        with ThreadPoolExecutor(max_workers=8) as pool:
            results = list(pool.map(lambda index: intake.parse_text(f"user-{index}", RESUME), range(16)))

        self.assertEqual(1, parser.calls)
        self.assertEqual(1, len({item.resume_hash for item in results}))
        self.assertEqual(16, len({item.profile.user_id for item in results}))

    def test_txt_and_docx_resume_input(self):
        self.assertEqual(RESUME.strip(), resume_text_from_bytes("resume.txt", RESUME.encode()))

        document = Document()
        document.add_paragraph("Software Engineer")
        document.add_paragraph("Python and FastAPI")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "PostgreSQL"
        table.cell(0, 1).text = "AWS"
        content = BytesIO()
        document.save(content)
        extracted = resume_text_from_bytes("resume.docx", content.getvalue())
        self.assertIn("Software Engineer", extracted)
        self.assertIn("Python and FastAPI", extracted)
        self.assertIn("PostgreSQL | AWS", extracted)

        with self.assertRaisesRegex(ValueError, "Supported resume formats"):
            resume_text_from_bytes("resume.pages", b"invalid")

    def test_common_alternative_section_headings_are_recognized(self):
        text = """WORK HISTORY
Backend Engineer
2020 - 2023
Used Python and FastAPI.
PERSONAL PROJECTS
Built a React application.
CORE COMPETENCIES
Python, FastAPI, React
ACADEMIC QUALIFICATIONS
Master's in Computer Science
CERTIFICATES
AWS Certified Developer
"""
        parsed = self.intake.parse_text("candidate-a", text)
        self.assertIn("fastapi", parsed.profile.professional_skills)
        self.assertIn("react", parsed.profile.project_skills)
        self.assertTrue(parsed.profile.education)
        self.assertTrue(parsed.profile.certifications)

    def test_resume_answers_flow_through_existing_retrieval_and_ranking(self):
        _, jobs, recommendations, intake = local_intake_stack(MatchingSettings(embedding_dimensions=64))
        jobs.process({
            "id": "backend", "job_title": "Backend Software Engineer", "company": "Acme",
            "description": "Build Python FastAPI APIs using PostgreSQL and AWS.",
            "skills_required": ["Python", "FastAPI", "PostgreSQL"],
            "skills_preferred": ["AWS"], "experience_min": 2, "experience_max": 6,
            "location": "Bengaluru, India", "work_mode": "Hybrid", "job_type": "Full Time",
        })
        jobs.process({
            "id": "frontend", "job_title": "Frontend Engineer", "company": "Web Co",
            "description": "Build React and TypeScript web interfaces.",
            "skills_required": ["React", "TypeScript"], "experience_min": 2,
            "location": "Bengaluru, India", "work_mode": "Hybrid", "job_type": "Full Time",
        })

        service = IntakeRecommendationService(intake, recommendations)
        response = service.recommend(ResumeRecommendationRequest(
            user_id="candidate-a", resume_text=RESUME, answers=required_answers(), limit=10,
        ))

        self.assertEqual("backend", response.recommendations.recommendations[0].job_id)
        self.assertEqual("candidate-a", response.intake.profile.user_id)
        self.assertTrue(response.recommendations.profile_hash)
        self.assertIn("python", response.recommendations.recommendations[0].matched_required_skills)


if __name__ == "__main__":
    unittest.main()
