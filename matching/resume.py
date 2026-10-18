from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
import threading
from typing import Any

from matching.hashing import canonical_employment_type, canonical_work_mode, normalize_text
from matching.schemas import CandidateIntakeResponse, CandidateProfile, EditableQuestion
from matching.taxonomy import Taxonomy

PARSER_VERSION = "resume-v2"
MAX_RESUME_BYTES = 10 * 1024 * 1024
MAX_RESUME_TEXT = 500_000
_SECTION = re.compile(
    r"(?im)^\s*(relevant experience|professional experience|work experience|experience|employment history|employment|work history|career history|"
    r"selected projects?|personal projects?|key projects?|projects?|education|academic qualifications?|skills?|technical skills|"
    r"core skills|core competencies|technical expertise|technologies|certifications?|certificates?|licenses?|executive summary|professional summary|summary)\s*:?[ \t]*$"
)
_DEGREE = re.compile(r"\b(?:bachelor(?:'s)?|master(?:'s)?|ph\.?d\.?|doctorate|associate(?:'s)?|b\.?tech|m\.?tech|b\.?e\.?|m\.?s\.?|b\.?s\.?)\b[^\n]{0,100}", re.I)
_CERTIFICATION = re.compile(r"\b(?:PMP|CISSP|CPA|CFA|AWS Certified[^\n,;]*|Azure Certified[^\n,;]*|Google Cloud Certified[^\n,;]*)\b", re.I)
_DATE_RANGE = re.compile(
    r"(?P<start_month>Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)?\s*"
    r"(?P<start_year>(?:19|20)\d{2})\s*(?:-|–|—|to)\s*"
    r"(?:(?P<end_month>Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)?\s*"
    r"(?P<end_year>(?:19|20)\d{2})|(?P<present>present|current|now))",
    re.I,
)
_MONTHS = {name.casefold(): index for index, names in enumerate(
    (("jan", "january"), ("feb", "february"), ("mar", "march"), ("apr", "april"),
     ("may",), ("jun", "june"), ("jul", "july"), ("aug", "august"),
     ("sep", "september"), ("oct", "october"), ("nov", "november"), ("dec", "december")), 1)
    for name in names}
_EMAIL = re.compile(r"(?<![\w.+-])[\w.+-]+@[\w-]+(?:\.[\w-]+)+(?![\w.-])", re.I)
_LINKEDIN = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[^\s|,;]+", re.I)
_URL = re.compile(r"https?://[^\s|,;]+", re.I)
_PHONE = re.compile(r"(?<!\w)(?:\+\d{1,3}[\s().-]*)?(?:\d[\s().-]*){8,14}\d(?!\w)")


def resume_text_from_bytes(filename: str, content: bytes) -> str:
    if len(content) > MAX_RESUME_BYTES:
        raise ValueError("Resume files must be 10 MB or smaller")
    suffix = Path(filename or "resume.txt").suffix.casefold()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages).strip()
    if suffix == ".docx":
        from docx import Document
        document = Document(BytesIO(content))
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells if cell.text.strip()))
        for section in document.sections:
            lines.extend(paragraph.text for paragraph in section.header.paragraphs)
            lines.extend(paragraph.text for paragraph in section.footer.paragraphs)
        return "\n".join(line for line in lines if line.strip()).strip()
    if suffix in {".txt", ".md", ".rtf", ""}:
        return content.decode("utf-8", errors="replace").strip()
    raise ValueError("Supported resume formats are PDF, DOCX, TXT, MD and pasted text")


class ResumeParser:
    def __init__(self, taxonomy: Taxonomy, now: datetime | None = None) -> None:
        self.taxonomy = taxonomy
        self.now = now or datetime.now()

    def parse(self, user_id: str, text: str) -> tuple[CandidateProfile, dict[str, list[str]]]:
        cleaned = _clean(text)
        contact = _contact_details(cleaned)
        sections = _sections(cleaned)
        experience = sections.get("experience", "")
        projects = sections.get("projects", "")
        skills_section = sections.get("skills", "")
        all_skills = self.taxonomy.extract_skills(cleaned)
        professional_skills = self.taxonomy.extract_skills(experience)
        project_skills = [skill for skill in self.taxonomy.extract_skills(projects) if skill not in professional_skills]
        explicit_skills = self.taxonomy.extract_skills(skills_section)
        roles, families, role_evidence = _roles(experience, self.taxonomy)
        professional_years, internship_years, date_evidence = _experience_years(experience, self.now)
        education = list(dict.fromkeys(match.group().strip() for match in _DEGREE.finditer(sections.get("education", cleaned))))
        certifications = list(dict.fromkeys(match.group().strip() for match in _CERTIFICATION.finditer(sections.get("certifications", cleaned))))
        career = _career_level(professional_years)
        evidence = {
            "skills": _evidence_lines(cleaned, all_skills),
            "professional_skills": _evidence_lines(experience, professional_skills),
            "project_skills": _evidence_lines(projects, project_skills),
            "roles": role_evidence,
            "experience_dates": date_evidence,
            "education": education,
            "certifications": certifications,
        }
        profile = CandidateProfile(
            user_id=user_id,
            **contact,
            skills=list(dict.fromkeys(explicit_skills + all_skills)),
            professional_skills=professional_skills,
            project_skills=project_skills,
            roles=roles,
            professional_roles=roles,
            role_families=families,
            professional_years=professional_years,
            internship_years=internship_years,
            education=education,
            certifications=certifications,
            career_level=career,
            processing_version=PARSER_VERSION,
        )
        return profile, evidence


def _contact_details(text: str) -> dict[str, Any]:
    """Extract conservative contact fields from the resume header."""
    first_section = _SECTION.search(text)
    header = text[:first_section.start() if first_section else min(len(text), 1500)]
    lines = [line.strip() for line in header.splitlines() if line.strip()][:15]
    email_match = _EMAIL.search(header)
    linkedin_match = _LINKEDIN.search(header)
    phone_match = _PHONE.search(header)
    full_name = None
    for line in lines:
        if any(token in line.casefold() for token in ("resume", "curriculum vitae", "linkedin", "github", "http", "@")):
            continue
        if any(char.isdigit() for char in line) or len(line) > 70:
            continue
        words = re.findall(r"[A-Za-z][A-Za-z'.-]*", line)
        if 2 <= len(words) <= 5 and len(" ".join(words)) >= 4:
            full_name = " ".join(words)
            break
    linkedin = linkedin_match.group().rstrip(".)") if linkedin_match else None
    urls = [match.group().rstrip(".)") for match in _URL.finditer(header)]
    websites = list(dict.fromkeys(url for url in urls if "linkedin.com/" not in url.casefold()))
    return {
        "full_name": full_name,
        "email": email_match.group() if email_match else None,
        "phone": phone_match.group().strip() if phone_match else None,
        "linkedin": linkedin,
        "websites": websites,
    }


class CandidateIntakeService:
    def __init__(self, parser: ResumeParser, questions_path: str | Path | None = None) -> None:
        self.parser = parser
        self.questions_path = Path(questions_path) if questions_path else Path(__file__).with_name("resume_questions.json")
        self._cache: dict[str, tuple[CandidateProfile, dict[str, list[str]]]] = {}
        self._key_locks: dict[str, threading.Lock] = {}
        self._lock = threading.RLock()

    def parse_text(self, user_id: str, text: str) -> CandidateIntakeResponse:
        cleaned = _clean(text)
        if not cleaned:
            raise ValueError("Resume text is empty")
        if len(cleaned) > MAX_RESUME_TEXT:
            raise ValueError("Resume text is too long")
        resume_hash = sha256(cleaned.encode()).hexdigest()
        key = f"{PARSER_VERSION}:{self.parser.taxonomy.version}:{resume_hash}"
        with self._lock:
            cached = self._cache.get(key)
            key_lock = self._key_locks.setdefault(key, threading.Lock())
        if cached is None:
            with key_lock:
                with self._lock:
                    cached = self._cache.get(key)
                if cached is None:
                    cached = self.parser.parse(user_id, cleaned)
                    with self._lock:
                        self._cache[key] = cached
        profile, evidence = deepcopy(cached)
        profile.user_id = user_id
        questions = self._questions(profile, evidence)
        return CandidateIntakeResponse(user_id=user_id, resume_hash=resume_hash, parser_version=PARSER_VERSION,
                                       profile=profile, questions=questions, evidence=evidence)

    def apply_answers(self, intake: CandidateIntakeResponse, answers: dict[str, Any]) -> CandidateIntakeResponse:
        profile = intake.profile.model_copy(deep=True)
        merged = {question.id: answers.get(question.id, question.answer) for question in intake.questions}
        required = [question.label for question in intake.questions
                    if question.required and ((question.id not in answers and not question.confirmed)
                                              or _empty(merged.get(question.id)))]
        if required:
            raise ValueError("Required candidate questions must be confirmed: " + ", ".join(required))
        for question in intake.questions:
            if question.id in answers:
                _validate_answer(question, answers[question.id])
        if "target_roles" in merged: profile.roles = _list(merged["target_roles"])
        if "confirmed_skills" in merged:
            confirmed = self.parser.taxonomy.normalize_skills(_list(merged["confirmed_skills"]))
            profile.professional_skills = [x for x in profile.professional_skills if x in confirmed]
            profile.project_skills = [x for x in profile.project_skills if x in confirmed]
            profile.skills = confirmed
        if not _empty(merged.get("professional_years")): profile.professional_years = max(0.0, float(merged["professional_years"]))
        if "preferred_locations" in merged: profile.locations = _list(merged["preferred_locations"])
        if "work_modes" in merged:
            profile.work_mode_preferences = [canonical_work_mode(x) for x in _list(merged["work_modes"])
                                               if normalize_text(x) != "no preference"]
            profile.remote_preference = True if profile.work_mode_preferences == ["remote"] else None
        if "employment_types" in merged:
            profile.employment_preferences = [canonical_employment_type(x) for x in _list(merged["employment_types"])]
        if "hard_location_constraint" in merged:
            profile.hard_location_constraint = _boolean(merged["hard_location_constraint"])
        if "strict_role_constraint" in merged:
            profile.strict_role_constraint = _boolean(merged["strict_role_constraint"])
        if "strict_work_mode_constraint" in merged:
            profile.strict_work_mode_constraint = _boolean(merged["strict_work_mode_constraint"])
        if "strict_employment_constraint" in merged:
            profile.strict_employment_constraint = _boolean(merged["strict_employment_constraint"])
        if "strict_experience_constraint" in merged:
            profile.strict_experience_constraint = _boolean(merged["strict_experience_constraint"])
        if "minimum_salary" in merged:
            profile.minimum_salary = (None if _empty(merged["minimum_salary"])
                                      else max(0.0, float(merged["minimum_salary"])))
        if "salary_currency" in merged:
            profile.salary_currency = (None if _empty(merged["salary_currency"])
                                       else str(merged["salary_currency"]).strip().upper())
        if "salary_flexible" in merged:
            profile.salary_flexible = _boolean(merged["salary_flexible"])
        if "work_authorization" in merged:
            profile.work_authorization = (None if _empty(merged["work_authorization"])
                                          else str(merged["work_authorization"]).strip())
        if "requires_sponsorship" in merged:
            profile.requires_sponsorship = (None if _empty(merged["requires_sponsorship"])
                                            else _boolean(merged["requires_sponsorship"]))
        if "willing_to_relocate" in merged:
            profile.willing_to_relocate = (None if _empty(merged["willing_to_relocate"])
                                           else _boolean(merged["willing_to_relocate"]))
        if "max_job_age_days" in merged:
            profile.max_job_age_days = (None if _empty(merged["max_job_age_days"])
                                        else max(1, int(float(merged["max_job_age_days"]))))
        if "blocked_companies" in merged:
            profile.blocked_companies = _list(merged["blocked_companies"])
        if "career_level" in merged:
            profile.career_level = None if _empty(merged["career_level"]) else str(merged["career_level"])
        questions = [question.model_copy(update={"answer": merged.get(question.id),
                                                  "confirmed": question.confirmed or question.id in answers})
                     for question in intake.questions]
        return intake.model_copy(update={"profile": profile, "questions": questions})

    def _questions(self, profile: CandidateProfile, evidence: dict[str, list[str]]) -> list[EditableQuestion]:
        definitions = json.loads(self.questions_path.read_text(encoding="utf-8"))
        defaults = {
            "target_roles": profile.roles,
            "confirmed_skills": profile.skills,
            "professional_years": profile.professional_years,
            "preferred_locations": profile.locations,
            "work_modes": profile.work_mode_preferences or ["no preference"],
            "employment_types": profile.employment_preferences or ["full time"],
            "hard_location_constraint": profile.hard_location_constraint,
            "strict_role_constraint": profile.strict_role_constraint,
            "strict_work_mode_constraint": profile.strict_work_mode_constraint,
            "strict_employment_constraint": profile.strict_employment_constraint,
            "strict_experience_constraint": profile.strict_experience_constraint,
            "minimum_salary": profile.minimum_salary,
            "salary_currency": profile.salary_currency,
            "salary_flexible": profile.salary_flexible,
            "work_authorization": profile.work_authorization,
            "requires_sponsorship": profile.requires_sponsorship,
            "willing_to_relocate": profile.willing_to_relocate,
            "max_job_age_days": profile.max_job_age_days or 30,
            "blocked_companies": profile.blocked_companies,
            "career_level": profile.career_level,
        }
        evidence_map = {"target_roles": evidence.get("roles", []), "confirmed_skills": evidence.get("skills", []),
                        "professional_years": evidence.get("experience_dates", [])}
        return [EditableQuestion.model_validate({**definition, "answer": defaults.get(definition["id"]),
                                                  "evidence": evidence_map.get(definition["id"], [])})
                for definition in definitions]


def _clean(text: str) -> str:
    value = str(text or "").replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"[ \t]+", " ", value)
    return re.sub(r"\n{3,}", "\n\n", value).strip()


def _sections(text: str) -> dict[str, str]:
    matches = list(_SECTION.finditer(text))
    aliases = {
        "relevant experience": "experience", "professional experience": "experience", "work experience": "experience", "experience": "experience",
        "employment history": "experience", "employment": "experience", "work history": "experience",
        "career history": "experience", "project": "projects", "projects": "projects",
        "selected project": "projects", "selected projects": "projects", "personal project": "projects",
        "personal projects": "projects", "key project": "projects", "key projects": "projects",
        "skill": "skills", "skills": "skills",
        "technical skills": "skills", "core skills": "skills", "core competencies": "skills",
        "technical expertise": "skills", "technologies": "skills", "certification": "certifications",
        "certifications": "certifications", "certificate": "certifications", "certificates": "certifications",
        "license": "certifications", "licenses": "certifications", "education": "education",
        "academic qualification": "education", "academic qualifications": "education",
        "executive summary": "summary", "professional summary": "summary", "summary": "summary",
    }
    result: dict[str, str] = {}
    for index, match in enumerate(matches):
        key = aliases.get(normalize_text(match.group(1)), "other")
        result[key] = result.get(key, "") + "\n" + text[match.end():matches[index+1].start() if index+1 < len(matches) else len(text)]
    return {key: value.strip() for key, value in result.items()}


def _roles(text: str, taxonomy: Taxonomy) -> tuple[list[str], list[str], list[str]]:
    roles, families, evidence = [], [], []
    for line in text.splitlines():
        family, subfamily = taxonomy.classify_role(line)
        if subfamily:
            roles.append(subfamily); evidence.append(line.strip())
        if family: families.append(family)
    return list(dict.fromkeys(roles)), list(dict.fromkeys(families)), list(dict.fromkeys(evidence))


def _experience_years(text: str, now: datetime) -> tuple[float | None, float | None, list[str]]:
    professional, internships, evidence = set(), set(), []
    for match in _DATE_RANGE.finditer(text):
        start = int(match["start_year"]) * 12 + _month(match["start_month"], 1)
        end = ((now.year * 12 + now.month) if match["present"] else int(match["end_year"]) * 12 + _month(match["end_month"], 12))
        if end < start or end - start > 720: continue
        line_start = text.rfind("\n", 0, match.start())
        line_start = 0 if line_start == -1 else line_start + 1
        line_end = text.find("\n", match.end())
        line_end = len(text) if line_end == -1 else line_end
        prev_line_start = text.rfind("\n", 0, line_start - 1)
        prev_line = text[max(0, prev_line_start + 1):line_start]
        context = f"{prev_line} {text[line_start:line_end]}"
        target = internships if re.search(r"\b(?:intern|internship|co-op)\b", context, re.I) else professional
        target.update(range(start, end + 1)); evidence.append(match.group())
    return (_years(professional), _years(internships), list(dict.fromkeys(evidence)))


def _month(value: str | None, default: int) -> int:
    return _MONTHS.get((value or "").casefold(), default)


def _years(months: set[int]) -> float | None:
    return round(len(months) / 12, 1) if months else None


def _career_level(years: float | None) -> str | None:
    if years is None: return None
    return "entry-level" if years < 2 else "mid" if years < 5 else "senior" if years < 9 else "lead"


def _evidence_lines(text: str, values: list[str]) -> list[str]:
    lines = []
    for line in text.splitlines():
        if any(re.search(r"(?<!\w)" + re.escape(value) + r"(?!\w)", line, re.I) for value in values): lines.append(line.strip())
    return list(dict.fromkeys(line for line in lines if line))


def _list(value: Any) -> list[str]:
    if value is None: return []
    if isinstance(value, str): return [item.strip() for item in value.split(",") if item.strip()]
    return list(dict.fromkeys(str(item).strip() for item in value if str(item).strip()))


def _empty(value: Any) -> bool:
    return value is None or value == "" or value == []


def _boolean(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    normalized = normalize_text(value)
    if normalized in {"true", "yes", "1"}:
        return True
    if normalized in {"false", "no", "0"}:
        return False
    raise ValueError(f"Expected a boolean answer, received {value!r}")


def _validate_answer(question: EditableQuestion, value: Any) -> None:
    if not question.required and _empty(value):
        return
    if question.kind.value == "number" and not _empty(value):
        try:
            float(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{question.label} must be a number") from exc
    if question.kind.value == "boolean":
        _boolean(value)
    if question.kind.value == "multi_select" and value is not None and not isinstance(value, (str, list, tuple, set)):
        raise ValueError(f"{question.label} must be a list")
    if question.options and not _empty(value):
        supplied = _list(value) if question.kind.value == "multi_select" else [str(value)]
        canonical = canonical_work_mode if question.id == "work_modes" else canonical_employment_type if question.id == "employment_types" else normalize_text
        allowed = {canonical(option) for option in question.options}
        invalid = [item for item in supplied if canonical(item) not in allowed]
        if invalid:
            raise ValueError(f"Invalid answer for {question.label}: {', '.join(invalid)}")
